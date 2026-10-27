"""
Market Report Generator - GUI Application
Using CustomTkinter for modern UI
"""

import customtkinter as ctk
from tkinter import filedialog, messagebox
import threading
import webbrowser
import os
import sys
from datetime import datetime

# Import our modules
from src.parser import parse_docx, MarketInput
from src.cache_manager import CacheManager
from src.ai_generator import AIGenerator
from src.payload_builder import build_payload, payload_to_dict
from src.renderer import render_report

# Configure CustomTkinter appearance
ctk.set_appearance_mode("dark")
ctk.set_default_color_theme("dark-blue")

# ── Design Tokens ──────────────────────────────────────────────────────────────
BG_APP        = "#0D1117"   # Window background
BG_CARD       = "#161B2E"   # Card / panel fill
BORDER_CARD   = "#2D3561"   # Card border
ACCENT        = "#6C63FF"   # Primary accent (electric indigo)
ACCENT_HOVER  = "#574FD6"   # Accent hover
BTN_SEC       = "#1E2A4A"   # Secondary button fill
BTN_SEC_HOVER = "#2D3D5C"   # Secondary button hover
SUCCESS       = "#00D48B"   # Success / active green
ERROR         = "#FF4D6D"   # Error red
RUNNING       = "#6C63FF"   # Running state (same as accent)
TEXT_PRIMARY  = "#F0F4FF"   # Primary text
TEXT_SECONDARY= "#8892A4"   # Muted / subtitle text
TEXT_TITLE    = "#A89CFF"   # Section header tint


class ProgressTracker:
    """Track and display progress for each stage."""

    def __init__(self, stages_frame):
        self.stages_frame = stages_frame
        self.stages = {}
        self.stage_widgets = {}

    def add_stage(self, name, description):
        """Add a new stage to track."""
        frame = ctk.CTkFrame(self.stages_frame, fg_color="transparent")
        frame.pack(fill="x", pady=6, padx=10)

        # Icon label
        icon_label = ctk.CTkLabel(
            frame, text="○", font=ctk.CTkFont(size=16), width=30,
            text_color="#4A5568"
        )
        icon_label.pack(side="left", padx=(0, 10))

        # Stage info
        info_frame = ctk.CTkFrame(frame, fg_color="transparent")
        info_frame.pack(side="left", fill="x", expand=True)

        name_label = ctk.CTkLabel(
            info_frame, text=name,
            font=ctk.CTkFont(size=12, weight="bold"),
            text_color=TEXT_PRIMARY,
            anchor="w",
        )
        name_label.pack(fill="x")

        desc_label = ctk.CTkLabel(
            info_frame,
            text=description,
            font=ctk.CTkFont(size=10),
            text_color=TEXT_SECONDARY,
            anchor="w",
        )
        desc_label.pack(fill="x")

        self.stages[name] = {
            "status": "pending",
            "description": description,
            "icon": icon_label,
            "name_label": name_label,
            "desc_label": desc_label,
        }
        self.stage_widgets[name] = frame

    def update(self, name, status, message=""):
        """Update stage status."""
        if name not in self.stages:
            return

        stage = self.stages[name]
        stage["status"] = status

        status_config = {
            "pending": ("○", "#4A5568",  "#4A5568"),
            "running": ("◉", RUNNING,    TEXT_SECONDARY),
            "success": ("✓", SUCCESS,    SUCCESS),
            "error":   ("✗", ERROR,      ERROR),
        }

        icon, icon_color, desc_color = status_config.get(
            status, status_config["pending"]
        )
        stage["icon"].configure(text=icon, text_color=icon_color)
        stage["name_label"].configure(text_color=icon_color if status != "pending" else TEXT_PRIMARY)
        if message:
            stage["desc_label"].configure(text=message, text_color=desc_color)

        self.stages_frame.update_idletasks()

    def reset(self):
        """Reset all stages to pending."""
        for name in self.stages:
            self.update(name, "pending")


class ReportGeneratorApp:
    """Main GUI Application."""

    def __init__(self):
        self.root = ctk.CTk()
        self.root.title("MarketIQ — Market Report Generator")
        self.root.geometry("960x720")
        self.root.minsize(860, 660)
        self.root.configure(fg_color=BG_APP)

        # Variables
        self.input_file = None
        self.output_file = None
        self.progress = None

        # Load config
        self.load_config()

        self.setup_ui()

    def load_config(self):
        """Load configuration."""
        import json

        config_path = os.path.join(os.path.dirname(__file__), "config.json")
        try:
            with open(config_path, "r") as f:
                self.config = json.load(f)
        except:
            self.config = {
                "openrouter_api_key": "",
                "model": "nvidia/nemotron-3-super-120b-a12b:free",
                "cache_ttl_hours": 24,
                "max_retries": 3,
                "template_file": "market_report_template_v2.docx",
            }

    def setup_ui(self):
        """Setup the user interface."""

        # Main container
        main_frame = ctk.CTkFrame(self.root, fg_color="transparent")
        main_frame.pack(fill="both", expand=True, padx=24, pady=20)

        # ── Branded Header ─────────────────────────────────────────────────────
        header_frame = ctk.CTkFrame(main_frame, fg_color="transparent")
        header_frame.pack(fill="x", pady=(0, 4))

        brand_dot = ctk.CTkLabel(
            header_frame,
            text="●",
            font=ctk.CTkFont(size=20),
            text_color=ACCENT,
            width=28,
        )
        brand_dot.pack(side="left", padx=(0, 8))

        title_label = ctk.CTkLabel(
            header_frame,
            text="MarketIQ",
            font=ctk.CTkFont(size=30, weight="bold"),
            text_color=TEXT_PRIMARY,
        )
        title_label.pack(side="left")

        # Logicawork branding — right side of header (clickable)
        logicawork_label = ctk.CTkLabel(
            header_frame,
            text="Powered by  Logicawork ↗",
            font=ctk.CTkFont(size=11),
            text_color=TEXT_SECONDARY,
            cursor="hand2",
        )
        logicawork_label.pack(side="right", padx=(0, 4))
        logicawork_label.bind(
            "<Button-1>",
            lambda e: webbrowser.open("https://logicawork.com/")
        )
        logicawork_label.bind(
            "<Enter>",
            lambda e: logicawork_label.configure(text_color=ACCENT)
        )
        logicawork_label.bind(
            "<Leave>",
            lambda e: logicawork_label.configure(text_color=TEXT_SECONDARY)
        )

        # Subtitle below header
        subtitle = ctk.CTkLabel(
            main_frame,
            text="AI-Powered Market Intelligence Report Generation",
            font=ctk.CTkFont(size=13),
            text_color=TEXT_SECONDARY,
        )
        subtitle.pack(anchor="w", pady=(0, 4))

        # Thin divider line
        divider = ctk.CTkFrame(main_frame, height=1, fg_color=BORDER_CARD)
        divider.pack(fill="x", pady=(0, 18))

        # ── File Selection Card ────────────────────────────────────────────────
        file_frame = ctk.CTkFrame(
            main_frame,
            fg_color=BG_CARD,
            border_width=1,
            border_color=BORDER_CARD,
            corner_radius=12,
        )
        file_frame.pack(fill="x", pady=(0, 16))

        file_inner = ctk.CTkFrame(file_frame, fg_color="transparent")
        file_inner.pack(pady=16, padx=18, fill="x")

        # Section label
        input_label = ctk.CTkLabel(
            file_inner,
            text="Input DOCX File",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color=TEXT_TITLE,
        )
        input_label.pack(anchor="w", pady=(0, 8))

        input_path_frame = ctk.CTkFrame(file_inner, fg_color="transparent")
        input_path_frame.pack(fill="x", pady=(0, 12))

        self.input_path_label = ctk.CTkLabel(
            input_path_frame,
            text="No file selected",
            font=ctk.CTkFont(size=11),
            text_color=TEXT_SECONDARY,
            anchor="w",
        )
        self.input_path_label.pack(side="left", fill="x", expand=True)

        browse_btn = ctk.CTkButton(
            input_path_frame,
            text="Browse",
            command=self.browse_input,
            width=100,
            fg_color=BTN_SEC,
            hover_color=BTN_SEC_HOVER,
            border_width=1,
            border_color=BORDER_CARD,
            corner_radius=8,
        )
        browse_btn.pack(side="right", padx=(10, 0))

        # Generate button
        self.generate_btn = ctk.CTkButton(
            file_inner,
            text="⚡  Generate Report",
            command=self.start_generation,
            height=44,
            font=ctk.CTkFont(size=14, weight="bold"),
            fg_color=ACCENT,
            hover_color=ACCENT_HOVER,
            corner_radius=10,
        )
        self.generate_btn.pack(fill="x", pady=(0, 0))

        # ── Progress Card ──────────────────────────────────────────────────────
        progress_frame = ctk.CTkFrame(
            main_frame,
            fg_color=BG_CARD,
            border_width=1,
            border_color=BORDER_CARD,
            corner_radius=12,
        )
        progress_frame.pack(fill="both", expand=True, pady=(0, 16))

        progress_inner = ctk.CTkFrame(progress_frame, fg_color="transparent")
        progress_inner.pack(pady=16, padx=18, fill="both", expand=True)

        progress_title = ctk.CTkLabel(
            progress_inner,
            text="Generation Progress",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color=TEXT_TITLE,
        )
        progress_title.pack(anchor="w", pady=(0, 10))

        # Scrollable progress stages
        stages_scroll = ctk.CTkScrollableFrame(
            progress_inner,
            fg_color=BG_APP,
            scrollbar_button_color=BORDER_CARD,
            scrollbar_button_hover_color=ACCENT,
            height=280,
            corner_radius=8,
        )
        stages_scroll.pack(fill="both", expand=True)

        self.progress_tracker = ProgressTracker(stages_scroll)

        # Define stages
        self.stages = [
            ("parsing",    "Parsing Input DOCX",    "Extracting market name and data..."),
            ("cache",      "Checking Cache",         "Looking for cached content..."),
            ("generating", "Generating Content",     "AI is creating report content..."),
            ("building",   "Building Payload",       "Preparing data for rendering..."),
            ("rendering",  "Rendering DOCX",         "Creating final report document..."),
            ("saving",     "Saving Output",          "Writing files to disk..."),
        ]

        for stage_id, name, desc in self.stages:
            self.progress_tracker.add_stage(stage_id, desc)

        # ── Output Card ────────────────────────────────────────────────────────
        output_frame = ctk.CTkFrame(
            main_frame,
            fg_color=BG_CARD,
            border_width=1,
            border_color=BORDER_CARD,
            corner_radius=12,
        )
        output_frame.pack(fill="x")

        output_inner = ctk.CTkFrame(output_frame, fg_color="transparent")
        output_inner.pack(pady=16, padx=18, fill="x")

        output_title = ctk.CTkLabel(
            output_inner,
            text="Output",
            font=ctk.CTkFont(size=13, weight="bold"),
            text_color=TEXT_TITLE,
        )
        output_title.pack(anchor="w", pady=(0, 10))

        # Output buttons
        output_btn_frame = ctk.CTkFrame(output_inner, fg_color="transparent")
        output_btn_frame.pack(fill="x")

        self.open_btn = ctk.CTkButton(
            output_btn_frame,
            text="Open Report",
            command=self.open_report,
            width=140,
            state="disabled",
            fg_color=BTN_SEC,
            hover_color=BTN_SEC_HOVER,
            border_width=1,
            border_color=BORDER_CARD,
            corner_radius=8,
        )
        self.open_btn.pack(side="left", padx=(0, 10))

        self.download_btn = ctk.CTkButton(
            output_btn_frame,
            text="Download",
            command=self.download_report,
            width=140,
            state="disabled",
            fg_color=BTN_SEC,
            hover_color=BTN_SEC_HOVER,
            border_width=1,
            border_color=BORDER_CARD,
            corner_radius=8,
        )
        self.download_btn.pack(side="left")

        self.output_path_label = ctk.CTkLabel(
            output_inner,
            text="No output generated yet",
            font=ctk.CTkFont(size=10),
            text_color=TEXT_SECONDARY,
            anchor="w",
        )
        self.output_path_label.pack(anchor="w", pady=(10, 0))

    def browse_input(self):
        """Browse for input DOCX file."""
        filename = filedialog.askopenfilename(
            title="Select Input DOCX",
            filetypes=[("DOCX files", "*.docx"), ("All files", "*.*")],
        )
        if filename:
            self.input_file = filename
            self.input_path_label.configure(
                text=os.path.basename(filename), text_color=SUCCESS
            )

    def start_generation(self):
        """Start report generation in a separate thread."""
        if not self.input_file:
            messagebox.showwarning("No File", "Please select an input DOCX file first.")
            return

        # Reset progress
        self.progress_tracker.reset()
        self.open_btn.configure(state="disabled")
        self.download_btn.configure(state="disabled")
        self.output_path_label.configure(text="Generating...", text_color=TEXT_SECONDARY)

        # Disable generate button
        self.generate_btn.configure(state="disabled", text="Generating...")

        # Start generation in thread
        thread = threading.Thread(target=self.generate_report)
        thread.daemon = True
        thread.start()

    def generate_report(self):
        """Generate the report (runs in separate thread)."""
        try:
            # Ensure input file is set
            if not self.input_file:
                self.update_progress("parsing", "error", "No input file selected")
                self.generation_complete(False)
                return

            # Stage 1: Parse
            self.update_progress("parsing", "running", "Parsing input DOCX...")
            market_input = parse_docx(self.input_file)

            if not market_input:
                self.update_progress("parsing", "error", "Failed to parse DOCX")
                self.generation_complete(False)
                return

            self.update_progress(
                "parsing", "success", f"Detected: {market_input.market_name}"
            )

            # Stage 2: Check cache
            self.update_progress("cache", "running", "Checking cache...")
            cache = CacheManager(
                cache_dir="cache", ttl_hours=self.config.get("cache_ttl_hours", 24)
            )
            cached_content = cache.get(
                market_input.market_name_upper, market_input.base_year
            )

            if cached_content:
                self.update_progress(
                    "cache", "success", "Using cached content (0 API calls)"
                )
                ai_content = cached_content
            else:
                self.update_progress(
                    "cache", "running", "Cache miss, generating new content..."
                )

                # Stage 3: Generate
                self.update_progress("generating", "running", "Calling AI API...")
                generator = AIGenerator(
                    api_key=self.config["openrouter_api_key"],
                    model=self.config["model"],
                    max_retries=self.config.get("max_retries", 3),
                    retry_delay=1,
                )

                ai_content = generator.generate(
                    market_input.market_name_upper, market_input.market_name_title
                )

                if ai_content:
                    cache.save(
                        market_input.market_name_upper,
                        ai_content,
                        market_input.base_year,
                    )
                    self.update_progress(
                        "generating", "success", "Content generated successfully"
                    )
                else:
                    self.update_progress(
                        "generating", "error", "API failed, using fallback"
                    )
                    from run import _get_fallback_content

                    ai_content = _get_fallback_content(market_input)
                    self.update_progress("cache", "success", "Using fallback content")

            # Stage 4: Build payload
            self.update_progress("building", "running", "Building payload...")
            payload = build_payload(ai_content, market_input)
            placeholder_dict = payload_to_dict(payload)
            self.update_progress(
                "building", "success", f"Ready: {len(placeholder_dict)} placeholders"
            )

            # Stage 5: Render
            self.update_progress("rendering", "running", "Rendering DOCX...")

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            safe_name = market_input.market_name.replace(" ", "_")
            output_docx = os.path.join("output", f"report_{safe_name}_{timestamp}.docx")
            output_json = os.path.join("output", f"report_{safe_name}_{timestamp}.json")

            template_path = self.config.get(
                "template_file", "market_report_template_v2.docx"
            )
            if not os.path.isabs(template_path):
                template_path = os.path.join(os.path.dirname(__file__), template_path)

            # Adjust cover title font size if title is long
            adjust_cover = market_input.cover_title_truncated

            success = render_report(
                template_path,
                placeholder_dict,
                output_docx,
                adjust_cover_title=adjust_cover,
            )

            if success:
                self.update_progress(
                    "rendering", "success", "DOCX rendered successfully"
                )
                self.output_file = output_docx

                # Stage 6: Save JSON
                self.update_progress("saving", "running", "Saving audit JSON...")
                import json

                with open(output_json, "w", encoding="utf-8") as f:
                    json.dump(
                        {
                            "market_name": market_input.market_name,
                            "generated_at": datetime.now().isoformat(),
                            "ai_content": ai_content,
                        },
                        f,
                        indent=2,
                        ensure_ascii=False,
                    )

                self.update_progress(
                    "saving", "success", f"Saved: {os.path.basename(output_json)}"
                )

                # Post-generation sanity check: ensure all placeholders were replaced
                try:
                    from docx import Document
                    doc = Document(output_docx)
                    # Gather all placeholder keys we attempted to replace
                    missing = []
                    for ph in placeholder_dict.keys():
                        # If the placeholder string still appears in any paragraph, it's missing
                        if any(ph in p.text for p in doc.paragraphs):
                            missing.append(ph)
                    if missing:
                        self.update_progress(
                            "rendering", "error",
                            f"Missing placeholders after render: {', '.join(missing)}"
                        )
                        self.generation_complete(False)
                        return
                except Exception as e:
                    self.update_progress("rendering", "error", f"Sanity check failed: {e}")
                    self.generation_complete(False)
                    return

                self.generation_complete(True)
            else:
                self.update_progress("rendering", "error", "Failed to render DOCX")
                self.generation_complete(False)

        except Exception as e:
            import traceback

            traceback.print_exc()
            messagebox.showerror("Error", f"Generation failed: {str(e)}")
            self.generation_complete(False)

    def update_progress(self, stage_id, status, message):
        """Update progress from thread (thread-safe)."""
        self.root.after(
            0, lambda: self.progress_tracker.update(stage_id, status, message)
        )

    def generation_complete(self, success):
        """Handle generation completion (thread-safe)."""
        self.root.after(0, lambda: self._complete_ui(success))

    def _complete_ui(self, success):
        """Update UI on completion."""
        self.generate_btn.configure(state="normal", text="⚡  Generate Report")

        if success:
            self.open_btn.configure(
                state="normal",
                fg_color=ACCENT,
                hover_color=ACCENT_HOVER,
                border_color=ACCENT,
            )
            self.download_btn.configure(
                state="normal",
                fg_color=ACCENT,
                hover_color=ACCENT_HOVER,
                border_color=ACCENT,
            )
            self.output_path_label.configure(
                text=f"Output: {self.output_file}", text_color=SUCCESS
            )
            messagebox.showinfo("Success", "Report generated successfully!")
        else:
            self.output_path_label.configure(
                text="Generation failed — check errors above", text_color=ERROR
            )

    def open_report(self):
        """Open the generated report."""
        if self.output_file and os.path.exists(self.output_file):
            os.startfile(self.output_file)

    def download_report(self):
        """Save copy of report to user-selected location."""
        if not self.output_file or not os.path.exists(self.output_file):
            return

        save_path = filedialog.asksaveasfilename(
            title="Save Report As",
            defaultextension=".docx",
            filetypes=[("DOCX files", "*.docx")],
            initialfile=os.path.basename(self.output_file),
        )

        if save_path:
            import shutil

            shutil.copy2(self.output_file, save_path)
            messagebox.showinfo("Saved", f"Report saved to:\n{save_path}")


def main():
    app = ReportGeneratorApp()
    app.root.mainloop()


if __name__ == "__main__":
    main()
